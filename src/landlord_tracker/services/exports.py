"""Local exports: formatted XLSX workbook (no CSV dumping).

Landlords and their accountants live in spreadsheets, so the export is a
first-class feature: styled header, frozen panes, sensible column widths,
currency formats and a dashboard summary sheet.
"""

from __future__ import annotations

import datetime as _dt
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..db import Database
from ..i18n import tr
from ..models import RentStatus
from . import calculations as calc
from .dashboard import build_dashboard

HEADER_FILL = PatternFill("solid", fgColor="1F3B63")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="1F3B63")
MONEY_FORMAT = '#,##0.00 "€"'
DATE_FORMAT = "yyyy-mm-dd"
THIN = Side(style="thin", color="D6DEEA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _write_sheet(ws, headers: list[str], rows: list[list], money_cols: set[int] = frozenset(),
                 date_cols: set[int] = frozenset()) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left")
        cell.border = BORDER
    ws.row_dimensions[1].height = 22

    for row in rows:
        ws.append(row)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = BORDER
            idx = cell.column
            if idx in money_cols:
                cell.number_format = MONEY_FORMAT
            elif idx in date_cols:
                cell.number_format = DATE_FORMAT

    for idx, header in enumerate(headers, start=1):
        longest = len(str(header))
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=idx, max_col=idx):
            for cell in row:
                longest = max(longest, len(str(cell.value or "")))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 3, 12), 46)

    ws.freeze_panes = "A2"


def _summary_sheet(wb: Workbook, db: Database, language: str) -> None:
    ws = wb.create_sheet(tr("dashboard"), 0)
    ws["A1"] = tr("app_name")
    ws["A1"].font = TITLE_FONT
    ws["A2"] = tr("export_generated", timestamp=_dt.datetime.now().strftime("%Y-%m-%d %H:%M"))
    ws["A2"].font = Font(italic=True, size=9, color="6B7A90")

    data = build_dashboard(db)
    ws["A4"] = tr("kpis")
    ws["A4"].font = TITLE_FONT
    row = 5
    for kpi in data.kpis:
        ws.cell(row=row, column=1, value=tr(kpi.key))
        cell = ws.cell(row=row, column=2, value=kpi.value)
        if kpi.kind == "money":
            cell.number_format = MONEY_FORMAT
        elif kpi.kind == "percent":
            cell.number_format = "0.0%"
        row += 1

    row += 1
    ws.cell(row=row, column=1, value=tr("renovations")).font = TITLE_FONT
    row += 1
    headers = [tr("renovation"), tr("property"), tr("budget"), tr("actual"),
               tr("variance"), tr("monthly_uplift"), tr("payback_months")]
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=idx, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    row += 1
    for reno in data.renovations:
        ws.cell(row=row, column=1, value=reno.title)
        ws.cell(row=row, column=2, value=reno.property_name)
        ws.cell(row=row, column=3, value=reno.budget).number_format = MONEY_FORMAT
        ws.cell(row=row, column=4, value=reno.actual).number_format = MONEY_FORMAT
        ws.cell(row=row, column=5, value=reno.variance).number_format = MONEY_FORMAT
        ws.cell(row=row, column=6, value=reno.monthly_uplift).number_format = MONEY_FORMAT
        ws.cell(row=row, column=7,
                value=(round(reno.payback_months, 1) if reno.payback_months else None))
        row += 1

    row += 1
    ws.cell(row=row, column=1, value=tr("privacy_line")).font = Font(italic=True, size=9, color="4C7A34")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 26
    for col in "CDEFG":
        ws.column_dimensions[col].width = 16


def export_workbook(db: Database, destination: Path, language: str = "en") -> Path:
    """Write a complete, styled workbook of the landlord's data."""
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    props = db.properties()
    prop_names = {p["id"]: p["name"] for p in props}
    tenant_names = {t["id"]: t["full_name"] for t in db.tenants()}

    _write_sheet(
        wb.create_sheet(tr("properties")),
        [tr("name"), tr("address"), tr("city"), tr("postcode"), tr("property_type"),
         tr("purchase_price"), tr("purchase_date"), tr("current_value"),
         tr("monthly_fixed_costs"), tr("size_m2"), tr("notes"), tr("created_at")],
        [[p["name"], p["address"], p["city"], p["postcode"], p["property_type"],
          calc.to_float(p["purchase_price"]), p["purchase_date"],
          calc.to_float(p["current_value"]), calc.to_float(p["monthly_fixed_costs"]),
          calc.to_float(p["size_m2"]), p["notes"], p["created_at"]] for p in props],
        money_cols={6, 8, 9}, date_cols={7},
    )

    _write_sheet(
        wb.create_sheet(tr("tenants")),
        [tr("full_name"), tr("property"), tr("phone"), tr("email"), tr("tax_number"),
         tr("emergency_contact"), tr("notes")],
        [[t["full_name"], t.get("property_name") or "", t["phone"], t["email"],
          t["tax_number"], t["emergency_contact"], t["notes"]] for t in db.tenants()],
    )

    _write_sheet(
        wb.create_sheet(tr("leases")),
        [tr("property"), tr("tenant"), tr("start_date"), tr("end_date"), tr("monthly_rent"),
         tr("deposit"), tr("due_day"), tr("last_increase"), tr("status"), tr("notes")],
        [[l.get("property_name") or "", l.get("tenant_name") or "", l["start_date"],
          l["end_date"], calc.to_float(l["monthly_rent"]), calc.to_float(l["deposit_amount"]),
          l["due_day"], l["last_increase_date"], tr(f"lease_status_{l['status']}"), l["notes"]]
         for l in db.leases()],
        money_cols={5, 6}, date_cols={3, 4, 8},
    )

    _write_sheet(
        wb.create_sheet(tr("rent")),
        [tr("period"), tr("property"), tr("tenant"), tr("due_amount"), tr("due_date"),
         tr("paid_amount"), tr("paid_date"), tr("balance"), tr("status"), tr("method"), tr("notes")],
        [[r["period"], r.get("property_name") or "", r.get("tenant_name") or "",
          calc.to_float(r["due_amount"]), r["due_date"], calc.to_float(r["paid_amount"]),
          r["paid_date"], calc.rent_balance(r["due_amount"], r["paid_amount"]),
          tr(f"rent_status_{r['status']}"), r["method"], r["notes"]]
         for r in db.rent_payments()],
        money_cols={4, 6, 8}, date_cols={5, 7},
    )

    _write_sheet(
        wb.create_sheet(tr("expenses")),
        [tr("date"), tr("property"), tr("category"), tr("amount"), tr("vendor"),
         tr("description"), tr("renovation"), tr("receipt")],
        [[e["date"], e.get("property_name") or "", tr(f"expense_cat_{e['category']}"),
          calc.to_float(e["amount"]), e["vendor"], e["description"],
          e.get("renovation_title") or "", e["receipt_path"]] for e in db.expenses()],
        money_cols={4}, date_cols={1},
    )

    _write_sheet(
        wb.create_sheet(tr("recurring")),
        [tr("property"), tr("category"), tr("amount"), tr("frequency"), tr("day_of_month"),
         tr("start_month"), tr("end_month"), tr("is_active"), tr("vendor"),
         tr("description"), tr("notes")],
        [[r.get("property_name") or "", tr(f"expense_cat_{r['category']}"),
          calc.to_float(r["amount"]), tr(f"freq_{r['frequency']}"),
          r["day_of_month"], r["start_month"], r["end_month"],
          tr("yes") if calc.as_flag(r.get("is_active", 1)) else tr("no"),
          r["vendor"], r["description"], r["notes"]]
         for r in db.recurring_expenses()],
        money_cols={3},
    )

    _write_sheet(
        wb.create_sheet(tr("renovations")),
        [tr("renovation"), tr("property"), tr("category"), tr("status"), tr("start_date"),
         tr("end_date"), tr("budget"), tr("logged_cost"), tr("variance"),
         tr("monthly_uplift"), tr("payback_months"), tr("contractor"), tr("notes")],
        [[r["title"], r.get("property_name") or "", tr(f"reno_cat_{r['category']}"),
          tr(f"reno_status_{r['status']}"), r["start_date"], r["end_date"],
          calc.to_float(r["budget"]), calc.to_float(r["logged_cost"]),
          calc.renovation_variance(r["budget"], r["logged_cost"]),
          calc.to_float(r["expected_monthly_rent_increase"]),
          (round(calc.renovation_payback_months(calc.to_float(r["logged_cost"]),
                                                calc.to_float(r["expected_monthly_rent_increase"])), 1)
           if calc.renovation_payback_months(calc.to_float(r["logged_cost"]),
                                             calc.to_float(r["expected_monthly_rent_increase"])) else None),
          r["contractor"], r["notes"]] for r in db.renovations()],
        money_cols={7, 8, 9, 10}, date_cols={5, 6},
    )

    _write_sheet(
        wb.create_sheet(tr("documents")),
        [tr("title"), tr("related_type"), tr("related_id"), tr("file_path"), tr("notes"),
         tr("created_at")],
        [[d["title"], d["related_type"], d["related_id"], d["file_path"], d["notes"],
          d["created_at"]] for d in db.documents()],
    )

    _summary_sheet(wb, db, language)
    wb.save(destination)
    return destination


def export_backup_zip_placeholder() -> None:  # pragma: no cover - documented no-op
    """Reserved for a future zipped backup (db + documents). See backups service."""
    raise NotImplementedError
