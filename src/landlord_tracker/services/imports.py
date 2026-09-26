"""Import helpers for user-owned Excel workbooks.

The importer is intentionally small and boring: it only accepts explicit sheets
and explicit column names, creates local rows, and never tries to guess hidden
meaning from free text. This keeps tenant data imports auditable.
"""

from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from ..db import Database
from ..models import Lease, Property, Tenant

PROPERTY_HEADERS = [
    "name", "address", "city", "postcode", "property_type",
    "purchase_price", "purchase_date", "current_value",
    "monthly_fixed_costs", "size_m2", "notes",
]
TENANT_HEADERS = [
    "full_name", "property_name", "phone", "email", "tax_number",
    "emergency_contact", "notes", "is_active",
]
LEASE_HEADERS = [
    "property_name", "tenant_full_name", "start_date", "end_date",
    "monthly_rent", "deposit_amount", "due_day", "last_increase_date",
    "status", "notes",
]

HEADER_FILL = PatternFill("solid", fgColor="1F3B63")
HEADER_FONT = Font(color="FFFFFF", bold=True)


@dataclass
class ImportResult:
    properties: int = 0
    tenants: int = 0
    leases: int = 0
    skipped: int = 0

    @property
    def total(self) -> int:
        return self.properties + self.tenants + self.leases


def _str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _float(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, str):
        value = value.replace("€", "").replace(" ", "").replace(",", ".")
    return float(value)


def _int(value: Any, default: int = 0) -> int:
    if value in (None, ""):
        return default
    return int(float(value))


def _date(value: Any) -> str:
    if value in (None, ""):
        return ""
    if isinstance(value, _dt.datetime):
        return value.date().isoformat()
    if isinstance(value, _dt.date):
        return value.isoformat()
    text = str(value).strip()
    # Keep already-normalized YYYY-MM-DD dates as-is. Users can also type text;
    # invalid dates should remain visible rather than silently mutated.
    return text


def _bool_int(value: Any, default: int = 1) -> int:
    if value in (None, ""):
        return default
    text = str(value).strip().lower()
    if text in {"0", "false", "no", "n", "inactive", "não", "nao"}:
        return 0
    return 1


def _rows(ws) -> list[dict[str, Any]]:
    headers = [_str(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, Any]] = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if not any(_str(v) for v in raw):
            continue
        rows.append({headers[i]: raw[i] if i < len(raw) else None for i in range(len(headers))})
    return rows


def _property_map(db: Database) -> dict[str, int]:
    return {_str(p["name"]).lower(): int(p["id"]) for p in db.properties()}


def _tenant_map(db: Database) -> dict[str, int]:
    return {_str(t["full_name"]).lower(): int(t["id"]) for t in db.tenants()}


def import_workbook(db: Database, path: Path) -> ImportResult:
    """Import Properties, Tenants and Leases sheets from an Excel workbook.

    Duplicate names are skipped rather than overwritten. This is safer for a
    desktop app with no undo stack: a user can inspect what changed, then edit
    rows manually if needed.
    """
    path = Path(path)
    wb = load_workbook(path, data_only=True)
    result = ImportResult()
    errors: list[str] = []

    if "Properties" in wb.sheetnames:
        existing = _property_map(db)
        for idx, row in enumerate(_rows(wb["Properties"]), start=2):
            name = _str(row.get("name"))
            if not name:
                errors.append(f"Properties row {idx}: name is required")
                continue
            if name.lower() in existing:
                result.skipped += 1
                continue
            prop_id = db.add_property(Property(
                name=name,
                address=_str(row.get("address")),
                city=_str(row.get("city")),
                postcode=_str(row.get("postcode")),
                property_type=_str(row.get("property_type")),
                purchase_price=_float(row.get("purchase_price")),
                purchase_date=_date(row.get("purchase_date")),
                current_value=_float(row.get("current_value")),
                monthly_fixed_costs=_float(row.get("monthly_fixed_costs")),
                size_m2=_float(row.get("size_m2")),
                notes=_str(row.get("notes")),
            ))
            existing[name.lower()] = prop_id
            result.properties += 1

    if "Tenants" in wb.sheetnames:
        properties = _property_map(db)
        existing = _tenant_map(db)
        for idx, row in enumerate(_rows(wb["Tenants"]), start=2):
            full_name = _str(row.get("full_name"))
            if not full_name:
                errors.append(f"Tenants row {idx}: full_name is required")
                continue
            if full_name.lower() in existing:
                result.skipped += 1
                continue
            property_name = _str(row.get("property_name"))
            property_id = properties.get(property_name.lower()) if property_name else None
            if property_name and property_id is None:
                errors.append(f"Tenants row {idx}: unknown property_name '{property_name}'")
                continue
            tenant_id = db.add_tenant(Tenant(
                full_name=full_name,
                property_id=property_id,
                phone=_str(row.get("phone")),
                email=_str(row.get("email")),
                tax_number=_str(row.get("tax_number")),
                emergency_contact=_str(row.get("emergency_contact")),
                notes=_str(row.get("notes")),
                is_active=_bool_int(row.get("is_active"), 1),
            ))
            existing[full_name.lower()] = tenant_id
            result.tenants += 1

    if "Leases" in wb.sheetnames:
        properties = _property_map(db)
        tenants = _tenant_map(db)
        for idx, row in enumerate(_rows(wb["Leases"]), start=2):
            property_name = _str(row.get("property_name"))
            tenant_name = _str(row.get("tenant_full_name"))
            start_date = _date(row.get("start_date"))
            if not property_name or not tenant_name or not start_date:
                errors.append(f"Leases row {idx}: property_name, tenant_full_name and start_date are required")
                continue
            property_id = properties.get(property_name.lower())
            tenant_id = tenants.get(tenant_name.lower())
            if property_id is None:
                errors.append(f"Leases row {idx}: unknown property_name '{property_name}'")
                continue
            if tenant_id is None:
                errors.append(f"Leases row {idx}: unknown tenant_full_name '{tenant_name}'")
                continue
            db.add_lease(Lease(
                property_id=property_id,
                tenant_id=tenant_id,
                start_date=start_date,
                end_date=_date(row.get("end_date")),
                monthly_rent=_float(row.get("monthly_rent")),
                deposit_amount=_float(row.get("deposit_amount")),
                due_day=_int(row.get("due_day"), 1),
                last_increase_date=_date(row.get("last_increase_date")),
                status=_str(row.get("status")) or "active",
                notes=_str(row.get("notes")),
            ))
            result.leases += 1

    if errors:
        preview = "\n".join(errors[:12])
        if len(errors) > 12:
            preview += f"\n... and {len(errors) - 12} more"
        raise ValueError(preview)
    return result


def _write_template_sheet(wb: Workbook, title: str, headers: list[str]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A2"
    for idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = min(max(len(header) + 4, 14), 28)


def create_import_template(destination: Path) -> Path:
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    rows = [
        ["Landlord Tracker import workbook"],
        ["Fill these sheets: Properties, Tenants, Leases."],
        ["Required: Properties.name; Tenants.full_name; Leases.property_name + tenant_full_name + start_date."],
        ["Dates: YYYY-MM-DD. Amounts: plain numbers like 750 or 750.50."],
        ["Names are used to link rows. property_name must match a Properties.name exactly."],
        ["Duplicate property or tenant names are skipped, not overwritten."],
        ["Remove the sample data inside Settings before importing your real data."],
    ]
    for row in rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=14, color="1F3B63")
    ws.column_dimensions["A"].width = 110

    _write_template_sheet(wb, "Properties", PROPERTY_HEADERS)
    _write_template_sheet(wb, "Tenants", TENANT_HEADERS)
    _write_template_sheet(wb, "Leases", LEASE_HEADERS)
    wb.save(destination)
    return destination
