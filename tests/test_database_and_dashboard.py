"""Database + dashboard + export tests against a real local SQLite file."""

from __future__ import annotations

import datetime as _dt
from pathlib import Path

import pytest

from landlord_tracker.db import Database
from landlord_tracker.models import Expense, Lease, Property, Renovation, RentPayment, Tenant
from landlord_tracker.services import calculations as calc
from landlord_tracker.services.demo import clear_demo, demo_summary, is_demo, load_demo
from landlord_tracker.services.dashboard import build_dashboard
from landlord_tracker.services.exports import export_workbook


@pytest.fixture()
def db(tmp_path: Path) -> Database:
    database = Database(data_dir=tmp_path)
    yield database
    database.close()


@pytest.fixture()
def seeded(db: Database) -> Database:
    load_demo(db)
    return db


# --------------------------------------------------------------------------
# schema / persistence
# --------------------------------------------------------------------------
def test_schema_creates_expected_tables(db: Database):
    names = {r["name"] for r in db.query(
        "SELECT name FROM sqlite_master WHERE type = 'table'")}
    for table in ("properties", "tenants", "leases", "rent_payments",
                  "expenses", "renovations", "documents", "settings"):
        assert table in names


def test_fresh_database_is_empty(db: Database):
    assert db.is_empty() is True
    assert db.count("properties") == 0


def test_integrity_check_passes(db: Database):
    assert db.integrity_check() == "ok"


def test_settings_round_trip(db: Database):
    assert db.get_setting("currency", "EUR") == "EUR"
    db.set_setting("currency", "USD")
    assert db.get_setting("currency") == "USD"
    db.set_setting("currency", "EUR")
    assert db.get_setting("currency") == "EUR"


def test_property_crud_and_persistence(tmp_path: Path):
    database = Database(data_dir=tmp_path)
    prop_id = database.add_property(Property(name="Test flat", city="Santarém",
                                             purchase_price=100000, current_value=140000))
    record = database.property(prop_id)
    assert record["name"] == "Test flat"
    database.update_property(prop_id, Property(name="Renamed", city="Lisboa",
                                               purchase_price=100000, current_value=150000))
    assert database.property(prop_id)["name"] == "Renamed"
    database.close()

    reopened = Database(data_dir=tmp_path)
    assert reopened.property(prop_id)["current_value"] == 150000
    reopened.delete_property(prop_id)
    assert reopened.property(prop_id) is None
    reopened.close()


def test_rent_status_is_persisted_and_drives_arrears(seeded: Database):
    rows = seeded.rent_payments()
    assert rows, "demo data should generate rent rows"
    statuses = {r["status"] for r in rows}
    assert "paid" in statuses
    assert calc.arrears(rows) > 0  # demo history intentionally includes friction


def test_rent_exists_prevents_duplicate_monthly_rows(seeded: Database):
    lease = seeded.active_leases()[0]
    period = seeded.rent_periods()[0]
    assert seeded.rent_exists(lease["id"], period) is not None


def test_document_vault_copies_file(db: Database, tmp_path: Path):
    source = tmp_path / "contract.txt"
    source.write_text("lease agreement", encoding="utf-8")
    stored = db.copy_document_into_vault(source)
    assert stored.exists()
    assert stored.parent == db.documents_dir
    assert source.exists(), "original file must never be moved"

    second = db.copy_document_into_vault(source)
    assert second.name != stored.name, "collisions must not overwrite"


def test_backup_and_prune(db: Database):
    db.add_property(Property(name="Backup me"))
    path = db.backup_to()
    assert path.exists() and path.stat().st_size > 0
    assert path in db.list_backups()
    assert db.prune_backups(keep=0) >= 1


# --------------------------------------------------------------------------
# demo data
# --------------------------------------------------------------------------
def test_demo_loads_and_is_removable(seeded: Database):
    summary = demo_summary(seeded)
    assert summary["properties"] == 3
    assert summary["tenants"] == 3
    assert summary["leases"] == 3
    assert summary["rent_rows"] >= 20
    assert summary["renovations"] == 3
    assert is_demo(seeded) is True

    clear_demo(seeded)
    assert seeded.is_empty() is True
    assert is_demo(seeded) is False


# --------------------------------------------------------------------------
# dashboard
# --------------------------------------------------------------------------
def test_dashboard_has_all_expected_kpis(seeded: Database):
    data = build_dashboard(seeded)
    keys = {kpi.key for kpi in data.kpis}
    for key in ("properties", "active_tenants", "occupancy", "rent_expected",
                "rent_collected", "rent_collection_rate", "arrears",
                "monthly_expenses", "net_cashflow", "portfolio_value",
                "purchase_total", "capital_gain", "renovation_invested",
                "annual_rent_roll", "gross_yield"):
        assert key in keys, f"missing KPI {key}"
    assert data.kpi("properties") == 3
    assert data.kpi("portfolio_value") == 168000 + 139000 + 124000


def test_dashboard_charts_are_populated(seeded: Database):
    data = build_dashboard(seeded, months=12)
    for name in ("cashflow", "rent_status", "expenses_by_category",
                 "portfolio_value", "purchase_vs_value", "property_ranking"):
        assert name in data.charts
    cashflow = data.charts["cashflow"]
    assert len(cashflow.labels) == 12
    assert len(cashflow.series["rent"]) == 12
    assert len(cashflow.series["expenses"]) == 12
    assert len(cashflow.series["net"]) == 12
    assert max(cashflow.series["rent"]) > 0


def test_dashboard_renovation_payback_is_computed_from_logged_cost(seeded: Database):
    data = build_dashboard(seeded)
    assert data.renovations
    bathroom = next(r for r in data.renovations if "Bathroom" in r.title)
    assert bathroom.actual == pytest.approx(1850 + 1420 + 780, abs=0.01)
    assert bathroom.budget == 4200
    assert bathroom.variance == pytest.approx(bathroom.actual - 4200, abs=0.01)
    assert bathroom.payback_months == pytest.approx(bathroom.actual / 75, abs=0.01)

    painting = next(r for r in data.renovations if "painting" in r.title.lower())
    assert painting.payback_months is None, "no rent uplift means no invented payback"


def test_dashboard_filtered_by_property(seeded: Database):
    prop = seeded.properties()[0]
    data = build_dashboard(seeded, property_id=prop["id"])
    assert data.kpi("properties") == 1
    assert data.kpi("portfolio_value") == prop["current_value"]


def test_dashboard_on_empty_database_gives_no_crash(db: Database):
    data = build_dashboard(db)
    assert data.kpi("properties") == 0
    assert data.kpi("net_cashflow") == 0
    assert data.charts["cashflow"].labels  # still returns a timeline


def test_dashboard_month_override_is_deterministic(seeded: Database):
    data = build_dashboard(seeded, today=_dt.date(2026, 6, 15), months=6)
    assert data.month == "2026-06"
    assert data.charts["cashflow"].labels == [
        "2026-01", "2026-02", "2026-03", "2026-04", "2026-05", "2026-06"]


# --------------------------------------------------------------------------
# export
# --------------------------------------------------------------------------
def test_export_workbook_contains_every_sheet(seeded: Database, tmp_path: Path):
    from openpyxl import load_workbook

    target = tmp_path / "out.xlsx"
    export_workbook(seeded, target)
    assert target.exists()
    wb = load_workbook(target)
    for sheet in ("Dashboard", "Properties", "Tenants", "Leases", "Rent",
                  "Expenses", "Renovations", "Documents"):
        assert sheet in wb.sheetnames
    assert wb["Properties"].max_row == 4  # header + 3 properties
    assert wb["Dashboard"]["A1"].value is not None


def test_export_works_on_empty_database(db: Database, tmp_path: Path):
    target = tmp_path / "empty.xlsx"
    export_workbook(db, target)
    assert target.exists()
