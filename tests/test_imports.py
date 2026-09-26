from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from landlord_tracker.db import Database
from landlord_tracker.services.imports import create_import_template, import_workbook


def test_create_import_template_has_expected_sheets(tmp_path: Path):
    path = create_import_template(tmp_path / "tenant-import-template.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == ["README", "Properties", "Tenants", "Leases"]
    assert wb["Properties"]["A1"].value == "name"
    assert wb["Tenants"]["A1"].value == "full_name"
    assert wb["Leases"]["A1"].value == "property_name"


def test_import_workbook_loads_properties_tenants_and_leases(tmp_path: Path):
    db = Database(data_dir=tmp_path / "data")
    workbook = tmp_path / "import.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Properties"
    ws.append(["name", "address", "city", "postcode", "property_type",
               "purchase_price", "purchase_date", "current_value",
               "monthly_fixed_costs", "size_m2", "notes"])
    ws.append(["Rua das Flores 12", "Rua das Flores 12", "Santarém", "2000-000",
               "Apartment", 100000, "2021-01-01", 125000, 55, 72, "real property"])

    ws = wb.create_sheet("Tenants")
    ws.append(["full_name", "property_name", "phone", "email", "tax_number",
               "emergency_contact", "notes", "is_active"])
    ws.append(["Ana Ferreira", "Rua das Flores 12", "+351 912 000 000",
               "ana@example.test", "123456789", "Rafa +351 999", "good payer", 1])

    ws = wb.create_sheet("Leases")
    ws.append(["property_name", "tenant_full_name", "start_date", "end_date",
               "monthly_rent", "deposit_amount", "due_day", "last_increase_date",
               "status", "notes"])
    ws.append(["Rua das Flores 12", "Ana Ferreira", "2024-01-01", "2024-12-31",
               680, 680, 5, "", "active", "imported lease"])
    wb.save(workbook)

    result = import_workbook(db, workbook)

    assert result.properties == 1
    assert result.tenants == 1
    assert result.leases == 1
    assert db.count("properties") == 1
    assert db.count("tenants") == 1
    assert db.count("leases") == 1
    assert db.tenants()[0]["property_name"] == "Rua das Flores 12"
    assert db.leases()[0]["monthly_rent"] == 680
    db.close()


def test_import_workbook_reports_unknown_property(tmp_path: Path):
    db = Database(data_dir=tmp_path / "data")
    workbook = tmp_path / "bad.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Tenants"
    ws.append(["full_name", "property_name"])
    ws.append(["Real Tenant", "Missing Property"])
    wb.save(workbook)

    with pytest.raises(ValueError, match="unknown property_name"):
        import_workbook(db, workbook)
    db.close()
